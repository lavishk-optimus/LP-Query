import logging
from typing import Optional, Union
from pathlib import Path
import io

try:
    from docx import Document
except ImportError:
    Document = None

try:
    import zipfile
    import xml.etree.ElementTree as ET
except ImportError:
    zipfile = None
    ET = None


logger = logging.getLogger(__name__)


class WordParserService:
    """Service for parsing and extracting text from Word documents"""
    
    def __init__(self):
        """Initialize the Word parser service"""
        self._check_dependencies()
        logger.info("WordParserService initialized successfully")
    
    def _check_dependencies(self):
        """Check if required dependencies are available"""
        if Document is None:
            logger.warning(
                "python-docx not found. Install with: pip install python-docx. "
                "Falling back to basic XML parsing."
            )
    
    def extract_text_from_word(self, file_input: Union[str, Path, bytes, io.BytesIO]) -> str:
        """
        Extract all text content from a Word document (.docx format)
        
        Args:
            file_input: Can be:
                - str/Path: File path to the Word document
                - bytes: Raw bytes of the Word document
                - io.BytesIO: BytesIO object containing the Word document
                
        Returns:
            str: Extracted text content from the document
            
        Raises:
            FileNotFoundError: If file path doesn't exist
            ValueError: If file format is not supported
            Exception: For other parsing errors
        """
        try:
            # Handle different input types
            if isinstance(file_input, (str, Path)):
                # File path input
                file_path = Path(file_input)
                if not file_path.exists():
                    raise FileNotFoundError(f"Word file not found: {file_path}")
                
                if not file_path.suffix.lower() in ['.docx']:
                    raise ValueError(f"Unsupported file format: {file_path.suffix}. Only .docx files are supported.")
                
                logger.info(f"Extracting text from Word file: {file_path}")
                return self._extract_from_file_path(file_path)
                
            elif isinstance(file_input, bytes):
                # Raw bytes input
                logger.info("Extracting text from Word document bytes")
                return self._extract_from_bytes(file_input)
                
            elif isinstance(file_input, io.BytesIO):
                # BytesIO input
                logger.info("Extracting text from Word document BytesIO")
                return self._extract_from_bytesio(file_input)
                
            else:
                raise ValueError(f"Unsupported input type: {type(file_input)}. Expected str, Path, bytes, or BytesIO.")
                
        except Exception as e:
            logger.error(f"Error extracting text from Word document: {str(e)}")
            raise
    
    def _extract_from_file_path(self, file_path: Path) -> str:
        """Extract text from Word file using file path"""
        if Document is not None:
            # Use python-docx if available
            try:
                doc = Document(file_path)
                return self._extract_text_with_docx(doc)
            except Exception as e:
                logger.warning(f"python-docx failed, falling back to XML parsing: {str(e)}")
                return self._extract_text_with_xml(file_path)
        else:
            # Fall back to XML parsing
            return self._extract_text_with_xml(file_path)
    
    def _extract_from_bytes(self, file_bytes: bytes) -> str:
        """Extract text from Word document bytes"""
        if Document is not None:
            # Use python-docx if available
            try:
                doc = Document(io.BytesIO(file_bytes))
                return self._extract_text_with_docx(doc)
            except Exception as e:
                logger.warning(f"python-docx failed, falling back to XML parsing: {str(e)}")
                return self._extract_text_with_xml_bytes(file_bytes)
        else:
            # Fall back to XML parsing
            return self._extract_text_with_xml_bytes(file_bytes)
    
    def _extract_from_bytesio(self, file_io: io.BytesIO) -> str:
        """Extract text from Word document BytesIO"""
        if Document is not None:
            # Use python-docx if available
            try:
                doc = Document(file_io)
                return self._extract_text_with_docx(doc)
            except Exception as e:
                logger.warning(f"python-docx failed, falling back to XML parsing: {str(e)}")
                file_io.seek(0)  # Reset position
                return self._extract_text_with_xml_bytes(file_io.read())
        else:
            # Fall back to XML parsing
            return self._extract_text_with_xml_bytes(file_io.read())
    
    def _extract_text_with_docx(self, doc) -> str:
        """Extract text using python-docx library"""
        try:
            text_parts = []
            
            # Extract text from paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text.strip())
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text_parts.append(cell.text.strip())
            
            extracted_text = '\n'.join(text_parts)
            logger.info(f"Successfully extracted {len(extracted_text)} characters using python-docx")
            return extracted_text
            
        except Exception as e:
            logger.error(f"Error extracting text with python-docx: {str(e)}")
            raise
    
    def _extract_text_with_xml(self, file_path: Path) -> str:
        """Extract text using direct XML parsing as fallback"""
        try:
            with zipfile.ZipFile(file_path, 'r') as docx_zip:
                return self._parse_document_xml(docx_zip)
        except Exception as e:
            logger.error(f"Error extracting text with XML parsing: {str(e)}")
            raise
    
    def _extract_text_with_xml_bytes(self, file_bytes: bytes) -> str:
        """Extract text using direct XML parsing from bytes"""
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes), 'r') as docx_zip:
                return self._parse_document_xml(docx_zip)
        except Exception as e:
            logger.error(f"Error extracting text with XML parsing from bytes: {str(e)}")
            raise
    
    def _parse_document_xml(self, docx_zip) -> str:
        """Parse the document.xml file to extract text"""
        try:
            # Read the main document XML
            document_xml = docx_zip.read('word/document.xml')
            root = ET.fromstring(document_xml)
            
            # Define namespace
            namespace = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
            
            # Extract all text elements
            text_parts = []
            for text_elem in root.findall('.//w:t', namespace):
                if text_elem.text:
                    text_parts.append(text_elem.text)
            
            extracted_text = ' '.join(text_parts)
            logger.info(f"Successfully extracted {len(extracted_text)} characters using XML parsing")
            return extracted_text
            
        except Exception as e:
            logger.error(f"Error parsing document XML: {str(e)}")
            raise
    
    def get_document_info(self, file_input: Union[str, Path, bytes, io.BytesIO]) -> dict:
        """
        Get basic information about the Word document
        
        Args:
            file_input: Word document input (same types as extract_text_from_word)
            
        Returns:
            dict: Document information including character count, word count, etc.
        """
        try:
            extracted_text = self.extract_text_from_word(file_input)
            
            info = {
                "character_count": len(extracted_text),
                "character_count_no_spaces": len(extracted_text.replace(' ', '')),
                "word_count": len(extracted_text.split()) if extracted_text else 0,
                "line_count": len(extracted_text.splitlines()) if extracted_text else 0,
                "is_empty": len(extracted_text.strip()) == 0
            }
            
            logger.info(f"Document info: {info}")
            return info
            
        except Exception as e:
            logger.error(f"Error getting document info: {str(e)}")
            raise